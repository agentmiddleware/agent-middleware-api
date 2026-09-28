"""Build an editable pilot workbook. Proposed economics are not actuals."""

from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
NAVY, TEAL = "15343E", "19796A"
MONEY = '$#,##0.00;($#,##0.00);"-"'
NUMBER = '#,##0.00;(#,##0.00);"-"'
PERCENT = '0.0%;(0.0%);"-"'
wb = Workbook()
wb.remove(wb.active)
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.properties.title = "One-tool pilot delivery tracker"


def sheet(name, title, subtitle, widths):
    s = wb.create_sheet(name)
    s.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        s.column_dimensions[chr(64 + i)].width = w
    for row, text in [(1, title), (2, subtitle)]:
        s.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=len(widths)
        )
        c = s.cell(row, 1, text)
        c.font = Font(
            name="Arial",
            size=18 if row == 1 else 10,
            bold=row == 1,
            color="FFFFFF" if row == 1 else NAVY,
        )
        c.fill = PatternFill("solid", fgColor=NAVY if row == 1 else "F1F6F4")
        s.row_dimensions[row].height = 34 if row == 1 else 40
    s.freeze_panes = "C7"
    s.page_setup.orientation = "landscape"
    s.page_setup.paperSize = s.PAPERSIZE_A4
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.oddFooter.center.text = name + " | Page &P of &N"
    return s


def inp(s, ref, value=None, fmt=None, note=None):
    c = s[ref]
    c.value = value
    c.font = Font(name="Arial", size=11, color="0000FF")
    c.fill = PatternFill("solid", fgColor="FFF2CC")
    c.border = Border(bottom=Side(style="hair", color="D7E1DD"))
    if fmt:
        c.number_format = fmt
    if note:
        c.comment = Comment(note, "Source / assumption")


def fx(s, ref, value, fmt=None):
    c = s[ref]
    c.value = value
    c.font = Font(name="Arial", size=11, color="008000" if "!" in value else "000000")
    if fmt:
        c.number_format = fmt


def dropdown(s, area, values):
    d = DataValidation(
        type="list", formula1='"' + ",".join(values) + '"', allow_blank=True
    )
    d.showErrorMessage = True
    d.error = "Choose a listed value."
    s.add_data_validation(d)
    d.add(area)


def headers(s, names):
    for i, name in enumerate(names, 1):
        c = s.cell(6, i, name)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
    s.row_dimensions[6].height = 34


def log(name, subtitle, names, widths, phase, dates):
    s = sheet(name, name + " · actual records", subtitle, widths)
    headers(s, names)
    for r in range(7, 107):
        for col in range(1, len(names) + 1):
            inp(s, f"{chr(64 + col)}{r}")
        for col in dates:
            s[f"{col}{r}"].number_format = "yyyy-mm-dd"
        s.row_dimensions[r].height = 27
    dropdown(s, f"{phase}7:{phase}106", ["Setup", "Recurring", "Acquisition", "Shared"])
    s.auto_filter.ref = f"A6:{chr(64 + len(names))}106"
    s.print_title_rows = "1:6"
    s.print_area = f"A1:{chr(64 + len(names))}20"
    return s


r = sheet(
    "Read me",
    "One-tool pilot · operating tracker",
    "September 8, 2026 | Internal draft | Use P-001; keep identities and payloads in the private customer system.",
    [29, 40, 42, 29],
)
notes = [
    (
        "Current state",
        "Qualification pending; no accepted order. Prices in Plan are hypotheses.",
    ),
    (
        "Edit cells",
        "Yellow/blue cells are inputs. Black is a formula; green references another sheet.",
    ),
    (
        "1 · Plan",
        "Review fee, labor allowance, provider costs and target margin. These do not become actuals.",
    ),
    (
        "2 · Logs",
        "Enter one engagement's Time, Costs, Revenue and Usage, using dates and the correct phase.",
    ),
    (
        "3 · Actuals",
        "Select inclusive dates and phase. Confirm each source complete only after reconciliation.",
    ),
    (
        "Unknown versus zero",
        "Unconfirmed logs produce Pending. Confirmed empty logs mean zero. Zero revenue/volume ratios are Undefined.",
    ),
    (
        "4 · Acceptance",
        "Pass needs partner-owned evidence, owner ID and date. Technical and commercial gates are separate.",
    ),
    ("5 · Example", "Fictional entries illustrate formats; never included in Actuals."),
    (
        "Cost boundary",
        "Time measures loaded labor. Costs excludes labor already counted in Time. Cash columns are nonlabor payments.",
    ),
    (
        "Provider allocation",
        "Record attributed usage and unabsorbed commitments once; do not double count included credits.",
    ),
    (
        "Revenue boundary",
        "Separate earned fee and cash dates/amounts; exclude pass-through tax; credits/refunds are negative adjustments.",
    ),
    (
        "Scope boundary",
        "One engagement per workbook. Rows 7–106 are included. Extend formulas if adding records beyond row 106.",
    ),
    (
        "Evidence",
        "Use private source references, not payloads or credentials. Keep the filled copy in approved storage.",
    ),
    (
        "Interpretation",
        "Operating analysis, not a general ledger. A planning reserve is not an actual expense.",
    ),
    ("Source", "PILOT_PACKAGE.md; ../unit-economics-reconsidered-2026-09-05/REPORT.md"),
]
for row, (label, value) in enumerate(notes, 4):
    r.cell(row, 1, label)
    r.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    r.cell(row, 2, value)
    r.row_dimensions[row].height = 38
r.print_area = "A1:D18"
r.freeze_panes = "A4"

p = sheet(
    "Plan",
    "Proposed economics · editable hypotheses",
    "Not a quote, bill or workload benchmark. Required hosting commitment remains unknown; replace allowances before offering managed service.",
    [44, 20, 75],
)
inputs = [
    (
        3,
        "Setup fee",
        2500,
        MONEY,
        "Internal hypothesis; reconsidered report section 19.",
    ),
    (4, "Setup delivery-hour limit", 20, NUMBER, "Proposed fixed-scope budget."),
    (
        5,
        "Setup resource / other allowance",
        100,
        MONEY,
        "Hypothesis; unknown contract increments excluded.",
    ),
    (
        6,
        "Loaded delivery cost / hour",
        75,
        MONEY,
        "Economic labor cost assumption, including unpaid founder effort.",
    ),
    (
        7,
        "Payment rate",
        0.029,
        PERCENT,
        "Planning assumption from the report; replace with actual terms.",
    ),
    (
        8,
        "Fixed fee per collection",
        0.3,
        MONEY,
        "Assumes one positive setup/monthly collection.",
    ),
    (
        9,
        "Planning concession reserve",
        0.01,
        PERCENT,
        "Hypothesis; excluded from Actuals.",
    ),
    (
        10,
        "Recurring monthly fee",
        1500,
        MONEY,
        "Internal test price; requires buyer value and scoped obligation.",
    ),
    (
        11,
        "Routine hours / month",
        4,
        NUMBER,
        "Proposed routine scope; not unlimited incident support.",
    ),
    (
        12,
        "Expected exception hours / month",
        0.125,
        NUMBER,
        "Unverified 100K-attempt illustration from the report.",
    ),
    (
        13,
        "Monthly resource allowance",
        50,
        MONEY,
        "Assumed attributed provider resource and backup cost.",
    ),
    (
        14,
        "Other monthly delivery allowance",
        20,
        MONEY,
        "Customer-attributable services, not shared overhead.",
    ),
    (
        15,
        "Recurring contribution target",
        0.70,
        PERCENT,
        "Internal threshold, not an industry benchmark.",
    ),
]
for row, label, value, fmt, note in inputs:
    p.cell(row, 1, label)
    inp(p, f"B{row}", value, fmt, note)
    p.cell(row, 3, note)
    p.row_dimensions[row].height = 34
outputs = {
    19: (
        "Setup delivery cost",
        '=IF(COUNT(B3:B9)<7,"Pending",B4*B6+B5+B3*(B7+B9)+IF(B3>0,B8,0))',
        MONEY,
    ),
    20: ("Setup contribution", '=IF(ISNUMBER(B19),B3-B19,"Pending")', MONEY),
    21: (
        "Setup margin",
        '=IF(ISNUMBER(B20),IF(B3=0,"Undefined",B20/B3),"Pending")',
        PERCENT,
    ),
    24: (
        "Recurring delivery cost",
        '=IF(COUNT(B6:B14)<9,"Pending",(B11+B12)*B6+B13+B14+B10*(B7+B9)+IF(B10>0,B8,0))',
        MONEY,
    ),
    25: ("Recurring contribution", '=IF(ISNUMBER(B24),B10-B24,"Pending")', MONEY),
    26: (
        "Recurring margin",
        '=IF(ISNUMBER(B25),IF(B10=0,"Undefined",B25/B10),"Pending")',
        PERCENT,
    ),
    27: (
        "Fee needed for target",
        '=IF(COUNT(B6:B15)<10,"Pending",IF(1-B7-B9-B15<=0,"Unreachable",((B11+B12)*B6+B13+B14+B8)/(1-B7-B9-B15)))',
        MONEY,
    ),
}
for row, (label, value, fmt) in outputs.items():
    p.cell(row, 1, label)
    fx(p, f"B{row}", value, fmt)
    p.row_dimensions[row].height = 27
p.print_area = "A1:C27"
p.freeze_panes = "B3"

t = log(
    "Time",
    "Work date and phase. Acquisition/Shared are excluded from selected Setup/Recurring delivery.",
    [
        "Date",
        "Phase",
        "Activity",
        "Hours",
        "Loaded $/hour",
        "Labor cost ($)",
        "Source / note",
    ],
    [15, 16, 31, 12, 18, 18, 48],
    "B",
    ["A"],
)
for row in range(7, 107):
    for col in ["D", "E"]:
        t[f"{col}{row}"].number_format = NUMBER
    fx(t, f"F{row}", f'=IF(COUNT(D{row}:E{row})<2,"",D{row}*E{row})', MONEY)
c = log(
    "Costs",
    "Accrual date is the service-period date. Cash date may differ. Exclude labor already in Time; cash is nonlabor.",
    [
        "Accrual date",
        "Cash date",
        "Phase",
        "Category",
        "Delivery cost ($)",
        "Cash paid ($)",
        "Source / allocation basis",
    ],
    [16, 16, 16, 27, 21, 20, 48],
    "C",
    ["A", "B"],
)
rv = log(
    "Revenue",
    "Platform/setup fees only; exclude pass-through taxes. Separate earned and cash dates. Credits/refunds are negative rows.",
    [
        "Earned date",
        "Collection date",
        "Phase",
        "Earned fee ($)",
        "Cash collected ($)",
        "Invoice / adjustment reference",
        "Source / note",
    ],
    [16, 16, 16, 20, 23, 36, 40],
    "C",
    ["A", "B"],
)
u = log(
    "Usage",
    "Match period, phase, environment and prospect. Ingress requests and accepted identities are different units.",
    [
        "Date",
        "Phase",
        "Accepted identities",
        "Ingress requests",
        "Dispatches",
        "Replays",
        "Denials",
        "Human cases",
        "Logical bytes",
        "Scope / source",
    ],
    [15, 15, 20, 20, 17, 15, 15, 17, 20, 46],
    "B",
    ["A"],
)
for s, cols in [(c, ["E", "F"]), (rv, ["D", "E"])]:
    for row in range(7, 107):
        for col in cols:
            s[f"{col}{row}"].number_format = MONEY
for row in range(7, 107):
    for col in "CDEFGHI":
        u[f"{col}{row}"].number_format = '#,##0;(#,##0);"-"'

a = sheet(
    "Actuals",
    "Confirmed actuals · one engagement and period",
    "Pending until confirmed complete. Plan and Example never feed these totals. Source records must share the selected scope.",
    [46, 24, 68],
)
for row, label in [
    (3, "Prospect ID"),
    (4, "Period start (inclusive)"),
    (5, "Period end (inclusive)"),
    (6, "Delivery phase"),
    (8, "Time records complete?"),
    (9, "Cost records complete?"),
    (10, "Revenue records complete?"),
    (11, "Usage records complete?"),
]:
    a.cell(row, 1, label)
    inp(
        a,
        f"B{row}",
        "P-001" if row == 3 else "Setup" if row == 6 else None,
        "yyyy-mm-dd" if row in [4, 5] else None,
    )
dropdown(a, "B6", ["Setup", "Recurring"])
for row in [8, 9, 10, 11]:
    dropdown(a, f"B{row}", ["Yes", "No"])
for row, note in [
    (8, "Confirm dates, hours, rates and phase are complete."),
    (9, "Confirm provider allocation and all delivery costs reconcile."),
    (10, "Confirm earned fees and cash records reconcile."),
    (11, "Confirm period-matched operational counts."),
]:
    a.cell(row, 3, note)
base = 'AND(ISNUMBER($B$4),ISNUMBER($B$5),$B$4<=$B$5,OR($B$6="Setup",$B$6="Recurring"))'


def span(tab, col):
    return tab + "!$" + col + "$7:$" + col + "$106"


def total(tab, amount, date, phase, complete):
    return (
        "=IF(AND("
        + base
        + ",$B$"
        + str(complete)
        + '="Yes"),SUMIFS('
        + span(tab, amount)
        + ","
        + span(tab, date)
        + ',">="&$B$4,'
        + span(tab, date)
        + ',"<="&$B$5,'
        + span(tab, phase)
        + ',$B$6),"Pending")'
    )


actual = {
    12: (
        "Period record status",
        "=IF(AND(" + base + ',COUNTIF(B8:B11,"Yes")=4),"Ready","Pending")',
        None,
    ),
    14: ("Earned platform revenue", total("Revenue", "D", "A", "C", 10), MONEY),
    15: ("Loaded delivery labor", total("Time", "F", "A", "B", 8), MONEY),
    16: ("Nonlabor delivery cost", total("Costs", "E", "A", "C", 9), MONEY),
    17: ("Total delivery cost", '=IF(COUNT(B15:B16)=2,SUM(B15:B16),"Pending")', MONEY),
    18: (
        "Delivery contribution",
        '=IF(AND(ISNUMBER(B14),ISNUMBER(B17)),B14-B17,"Pending")',
        MONEY,
    ),
    19: (
        "Contribution margin",
        '=IF(ISNUMBER(B18),IF(B14=0,"Undefined",B18/B14),"Pending")',
        PERCENT,
    ),
    20: ("Accepted identities", total("Usage", "C", "A", "B", 11), '#,##0;(#,##0);"-"'),
    21: (
        "Delivery cost / 1,000 identities",
        '=IF(AND(ISNUMBER(B17),ISNUMBER(B20)),IF(B20=0,"Undefined",B17/B20*1000),"Pending")',
        MONEY,
    ),
    22: ("Cash collected this period", total("Revenue", "E", "B", "C", 10), MONEY),
    23: ("Logged nonlabor cash paid", total("Costs", "F", "B", "C", 9), MONEY),
    24: (
        "Receipts less nonlabor payments",
        '=IF(COUNT(B22:B23)=2,B22-B23,"Pending")',
        MONEY,
    ),
    25: ("Delivery hours", total("Time", "D", "A", "B", 8), NUMBER),
}
for row, (label, value, fmt) in actual.items():
    a.cell(row, 1, label)
    fx(a, f"B{row}", value, fmt)
    a.row_dimensions[row].height = 28
a["C17"] = "Includes actual payment fees entered in Costs; no hypothetical reserve."
a["C24"] = (
    "Not company cash profit: excludes payroll cash, acquisition, shared overhead and financing."
)
a["C25"] = (
    "All selected-phase hours; inspect activities to split routine and incident work."
)
a.print_area = "A1:C25"
a.freeze_panes = "B3"

ac = sheet(
    "Acceptance",
    "Partner acceptance · recorded evidence required",
    "A01–A10 technical; A11 commercial. Pass needs status, partner owner ID, date and evidence. All start Pending.",
    [10, 44, 16, 24, 16, 46, 40],
)
headers(
    ac,
    ["ID", "Case", "Status", "Partner owner ID", "Date", "Evidence reference", "Notes"],
)
cases = [
    "Ownership and environment",
    "Ordinary success",
    "Identical replay",
    "Worker restart / stable identity",
    "Conflicting payload",
    "Scope and budget denial",
    "Auth / permit boundary",
    "Post-dispatch uncertainty",
    "Offline verification / tamper",
    "Operating burden measured",
    "Value and commercial commitment",
]
for row, case in enumerate(cases, 7):
    ac.cell(row, 1, f"A{row - 6:02}")
    ac.cell(row, 2, case)
    for col in "CDEFG":
        inp(
            ac,
            f"{col}{row}",
            "Pending" if col == "C" else None,
            "yyyy-mm-dd" if col == "E" else None,
        )
    ac.row_dimensions[row].height = 45
dropdown(ac, "C7:C17", ["Pending", "Pass", "Fail"])
ac["A3"] = "Technical"
fx(
    ac,
    "B3",
    '=IF(COUNTIF(C7:C16,"Fail")>0,"Fail",IF(COUNTIFS(C7:C16,"Pass",D7:D16,"<>",E7:E16,">0",F7:F16,"<>")=10,"Pass","Pending"))',
)
ac["A4"] = "Commercial"
fx(
    ac,
    "B4",
    '=IF(C17="Fail","Declined",IF(AND(C17="Pass",D17<>"",ISNUMBER(E17),F17<>""),"Committed","Pending"))',
)
ac.print_area = "A1:G17"

ex = sheet(
    "Example",
    "Fictional examples · excluded from Actuals",
    "Formats only. Not P-001 activity, bills or forecasts. Do not copy these as evidence.",
    [21, 24, 34, 20, 20, 36],
)
headers(
    ex,
    [
        "Log",
        "Example date",
        "Activity / category",
        "Quantity / cost",
        "Rate / cash",
        "Explanation",
    ],
)
values = [
    [
        "Time",
        datetime(2026, 9, 8),
        "Recurring routine support",
        3,
        75,
        "Hours and loaded rate.",
    ],
    [
        "Time",
        datetime(2026, 9, 8),
        "Recurring exception work",
        0.125,
        75,
        "Not measured activity.",
    ],
    [
        "Costs",
        datetime(2026, 9, 8),
        "Resource allocation",
        50,
        50,
        "Accrual and cash can differ.",
    ],
    ["Costs", datetime(2026, 9, 8), "Other delivery", 20, 20, "Exclude labor in Time."],
    [
        "Costs",
        datetime(2026, 9, 8),
        "Payment fee",
        29.271,
        29.271,
        "Real fee from statement.",
    ],
    [
        "Revenue",
        datetime(2026, 9, 8),
        "Earned fee / collection",
        999,
        999,
        "Separate dates in Revenue.",
    ],
    [
        "Usage",
        datetime(2026, 9, 8),
        "Accepted identities",
        100000,
        None,
        "Matched period and scope.",
    ],
]
for row, vals in enumerate(values, 7):
    for col, v in enumerate(vals, 1):
        ex.cell(row, col, v)
    ex.cell(row, 2).number_format = "yyyy-mm-dd"
    ex.row_dimensions[row].height = 34
for row, label, f, fmt in [
    (15, "Labor", "=D7*E7+D8*E8", MONEY),
    (16, "Nonlabor", "=SUM(D9:D11)", MONEY),
    (17, "Contribution", "=D12-D15-D16", MONEY),
    (18, "Margin", '=IF(D12=0,"Undefined",D17/D12)', PERCENT),
]:
    ex.cell(row, 1, label)
    fx(ex, f"D{row}", f, fmt)
ex["A20"] = "Difference from Plan"
ex.merge_cells("B20:F20")
ex["B20"] = "No hypothetical 1% reserve is subtracted from fictional actuals."
ex.row_dimensions[20].height = 35
ex.print_area = "A1:F20"

for s in wb:
    for row in s:
        for cell in row:
            if cell.font.name != "Arial":
                cell.font = Font(name="Arial", size=10, color=NAVY)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    s.sheet_properties.tabColor = TEAL
for s, area in [(p, "B19:B27"), (a, "B14:B25")]:
    s.conditional_formatting.add(
        area,
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=PatternFill("solid", fgColor="FCE8E6"),
        ),
    )
wb.active = 0
wb.save(ROOT / "pilot-tracker.xlsx")
print("Created nine-sheet pilot-tracker.xlsx")
